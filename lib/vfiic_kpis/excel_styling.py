from __future__ import annotations

import re
from datetime import date, datetime, time
from typing import Any

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from vfiic_kpis.excel_theme import (
    ColumnFormatRule,
    ExcelTheme,
    first_matching_rule,
    header_fill_argb,
    header_font_argb,
    resolve_header_label,
    row_stripe_fill_argb,
    semantic_color_argb,
    title_fill_argb,
    title_font_argb,
)


def _alignment_for(align: str) -> Alignment | None:
    value = align.lower()
    if value in ("", "general", "auto"):
        return None
    if value not in ("left", "right", "center", "justify"):
        return None
    return Alignment(horizontal=value, vertical="center")


def _collect_table_names(workbook: Workbook) -> set[str]:
    names: set[str] = set()
    for worksheet in workbook.worksheets:
        for table in worksheet.tables.values():
            names.add(table.displayName)
    return names


def _allocate_table_display_name(workbook: Workbook, seed: str) -> str:
    raw = re.sub(r"[^A-Za-z0-9_]", "_", seed).strip("_") or "Tabla"
    if raw[0].isdigit():
        raw = "T_" + raw
    base = raw[:200]
    used = _collect_table_names(workbook)
    candidate = base
    counter = 1
    while candidate in used:
        suffix = f"_{counter}"
        candidate = (base[: 200 - len(suffix)] + suffix)[:255]
        counter += 1
    return candidate[:255]


def _maybe_coerce_large_number(cell: Any, rule: ColumnFormatRule) -> None:
    if not rule.coerce_large_number_to_text:
        return
    value = cell.value
    if isinstance(value, float) and value.is_integer():
        iv = int(value)
        if abs(iv) >= 10**12:
            cell.value = str(iv)
        return
    if isinstance(value, int) and abs(value) >= 10**12:
        cell.value = str(value)


def _apply_body_formats(
    ws: Worksheet,
    internal_columns: list[str],
    theme: ExcelTheme,
    data_row_start: int,
    data_row_end: int | None = None,
    column_offset: int = 0,
) -> None:
    last_row = data_row_end if data_row_end is not None else (ws.max_row or 0)
    if last_row < data_row_start or not internal_columns:
        return

    for slot, internal in enumerate(internal_columns):
        col_idx = column_offset + slot + 1
        rule = first_matching_rule(internal, theme.column_formats) if internal else None
        for row_idx in range(data_row_start, last_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if rule is not None:
                _maybe_coerce_large_number(cell, rule)
                cell.number_format = rule.number_format
                align = _alignment_for(rule.align)
                if align is not None:
                    cell.alignment = align
            else:
                value = cell.value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif isinstance(value, (datetime, date, time)):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif isinstance(value, str):
                    cell.alignment = Alignment(horizontal="left", vertical="center")


def _paint_header_row(
    ws: Worksheet,
    theme: ExcelTheme,
    columns: list[str],
    row_idx: int = 1,
    column_offset: int = 0,
    explicit_labels: list[str] | None = None,
) -> None:
    fill = PatternFill(fill_type="solid", fgColor=header_fill_argb(theme))
    font = Font(bold=theme.header.bold, color=header_font_argb(theme))
    for slot, internal in enumerate(columns):
        col_idx = column_offset + slot + 1
        cell = ws.cell(row=row_idx, column=col_idx)
        if explicit_labels is not None and slot < len(explicit_labels) and explicit_labels[slot]:
            cell.value = explicit_labels[slot]
        else:
            cell.value = resolve_header_label(internal, theme)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _set_column_widths(ws: Worksheet, theme: ExcelTheme) -> None:
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    if max_col < 1:
        return

    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        longest = 0
        for row_idx in range(1, max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            longest = max(longest, len(str(value)))
        width = min(
            max(theme.column_width.min_width, longest + theme.column_width.margin),
            theme.column_width.max_width,
        )
        ws.column_dimensions[letter].width = width


def _apply_table(ws: Worksheet, workbook: Workbook, theme: ExcelTheme, header_row: int = 1) -> None:
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    if max_row < header_row + 1 or max_col < 1:
        return

    ref = f"A{header_row}:{get_column_letter(max_col)}{max_row}"
    display = _allocate_table_display_name(workbook, ws.title or "Tabla")
    table = Table(displayName=display, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=theme.table.style,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=theme.table.show_row_stripes,
        showColumnStripes=theme.table.show_column_stripes,
    )
    ws.add_table(table)


def apply_sheet_theme(
    ws: Worksheet,
    workbook: Workbook,
    internal_columns: list[str],
    theme: ExcelTheme,
) -> None:
    """Apply theme styling to a pandas-written sheet (row 1 = internal headers)."""
    if not internal_columns:
        return

    max_col_ws = ws.max_column or 0
    if max_col_ws < 1:
        return

    columns = list(internal_columns)[:max_col_ws]

    _paint_header_row(ws, theme, columns)
    _apply_body_formats(ws, columns, theme, data_row_start=2)
    _set_column_widths(ws, theme)

    ws.sheet_view.showGridLines = theme.layout.show_grid_lines
    if theme.layout.freeze_panes:
        ws.freeze_panes = theme.layout.freeze_panes

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    will_add_table = max_row >= 2 and max_col >= 1
    if will_add_table:
        ws.auto_filter = AutoFilter()
    elif max_row >= 1 and max_col >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    _apply_table(ws, workbook, theme)
    _paint_header_row(ws, theme, columns)


def paint_title_band(
    ws: Worksheet,
    title: str,
    theme: ExcelTheme,
    row_idx: int,
    column_count: int,
) -> None:
    """Paint a merged title band on the given row."""
    if column_count < 1:
        return
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=column_count)
    cell = ws.cell(row=row_idx, column=1)
    cell.value = title
    cell.fill = PatternFill(fill_type="solid", fgColor=title_fill_argb(theme))
    cell.font = Font(bold=theme.title.bold, color=title_font_argb(theme), size=12)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row_idx].height = 24


def paint_block_header(
    ws: Worksheet,
    theme: ExcelTheme,
    internal_columns: list[str],
    explicit_labels: list[str],
    row_idx: int,
) -> None:
    """Paint headers using explicit labels (supports dynamic month labels)."""
    _paint_header_row(
        ws,
        theme,
        internal_columns,
        row_idx=row_idx,
        explicit_labels=explicit_labels,
    )


def paint_block_body(
    ws: Worksheet,
    theme: ExcelTheme,
    internal_columns: list[str],
    data_row_start: int,
    data_row_end: int,
) -> None:
    _apply_body_formats(
        ws,
        internal_columns,
        theme,
        data_row_start=data_row_start,
        data_row_end=data_row_end,
    )


def paint_block_data_row_stripes(
    ws: Worksheet,
    theme: ExcelTheme,
    *,
    data_row_start: int,
    data_row_end: int,
    column_count: int,
) -> None:
    """Alterna fondo blanco / azul claro en las filas de datos de un bloque del comparativo."""
    if not theme.row_stripes.enabled or column_count < 1:
        return
    if data_row_end < data_row_start:
        return
    for offset, row_idx in enumerate(range(data_row_start, data_row_end + 1)):
        fill = PatternFill(fill_type="solid", fgColor=row_stripe_fill_argb(theme, offset))
        for col_idx in range(1, column_count + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill


def paint_porcentaje_bold(
    ws: Worksheet,
    *,
    column_indices: list[int],
    data_row_start: int,
    data_row_end: int,
) -> None:
    """Aplica `bold=True` como estilo base a las celdas de las columnas indicadas.

    Pensado para columnas de porcentaje del comparativo apilado. El bold se
    aplica como estilo base (no via formato condicional) para que sobreviva
    la evaluación dinámica de color por valor: las reglas condicionales solo
    sobrescriben `font.color`, dejando el bold intacto.
    """
    if data_row_end < data_row_start or not column_indices:
        return
    for col_idx in column_indices:
        for row_idx in range(data_row_start, data_row_end + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            current = cell.font
            cell.font = Font(
                name=current.name,
                size=current.size,
                bold=True,
                italic=current.italic,
                underline=current.underline,
                strike=current.strike,
                color=current.color,
            )


def apply_semantic_conditional_format(
    ws: Worksheet,
    theme: ExcelTheme,
    *,
    ranges: list[str],
) -> None:
    """Registra reglas de formato condicional sobre los rangos indicados.

    Cada celda del rango se colorea en tiempo de evaluación según su valor
    numérico:

      - `> 0` → color `positive` del tema (azul por defecto).
      - `< 0` → color `negative` del tema (rojo por defecto).
      - `= 0` → color `neutral` del tema (negro por defecto).

    Como las reglas se evalúan en Excel al recalcular fórmulas, los colores
    se actualizan cuando el usuario edita los valores base que alimentan las
    fórmulas de diferencia y variación.
    """
    if not ranges:
        return

    positive_color = semantic_color_argb(theme, "positive")
    negative_color = semantic_color_argb(theme, "negative")
    neutral_color = semantic_color_argb(theme, "neutral")

    for rng in ranges:
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=positive_color)),
        )
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator="lessThan", formula=["0"], font=Font(color=negative_color)),
        )
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator="equal", formula=["0"], font=Font(color=neutral_color)),
        )


def finalize_uniform_widths(ws: Worksheet, theme: ExcelTheme) -> None:
    """Calcula anchos uniformes considerando el contenido completo de la hoja."""
    _set_column_widths(ws, theme)
