from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC = PROJECT_ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from openpyxl import Workbook

from vfiic_kpis.excel_styling import apply_sheet_theme
from vfiic_kpis.excel_theme import (
    first_matching_rule,
    friendly_header_label,
    load_excel_theme,
    match_column_rule,
    resolve_header_label,
)
from vfiic_kpis.paths import DEFAULT_COMPARISON_V2_THEME, DEFAULT_PARTITIONED_THEME


class TestExcelTheme(unittest.TestCase):
    def test_load_default_comparison_v2_theme(self) -> None:
        theme = load_excel_theme(DEFAULT_COMPARISON_V2_THEME)
        self.assertEqual(theme.header.fill_rgb, "1F4E79")
        self.assertEqual(theme.header.font_color_rgb, "FFFFFF")
        self.assertTrue(theme.header.bold)
        self.assertIn("diferencia_mom", theme.header_labels)
        self.assertIn("porcentaje_yoy", theme.header_labels)

    def test_load_default_partitioned_theme(self) -> None:
        theme = load_excel_theme(DEFAULT_PARTITIONED_THEME)
        self.assertTrue(theme.header_labels)
        self.assertTrue(theme.column_formats)

    def test_friendly_header_label(self) -> None:
        self.assertEqual(friendly_header_label("area_id"), "Area Id")

    def test_resolve_header_label_override(self) -> None:
        theme = load_excel_theme(DEFAULT_PARTITIONED_THEME)
        self.assertEqual(resolve_header_label("area_id", theme), "ID área")

    def test_first_matching_rule_for_v2_percentage(self) -> None:
        theme = load_excel_theme(DEFAULT_COMPARISON_V2_THEME)
        rule = first_matching_rule("porcentaje_mom", theme.column_formats)
        self.assertIsNotNone(rule)
        assert rule is not None
        self.assertEqual(rule.match, "exact")
        self.assertIn("%", rule.number_format)

    def test_regex_column_format(self) -> None:
        body = """
[[column_formats]]
match = "regex"
pattern = ".*_id$"
number_format = "@"
align = "left"
"""
        with tempfile.NamedTemporaryFile("w", suffix=".toml", delete=False, encoding="utf-8") as handle:
            handle.write(body)
            path = Path(handle.name)
        try:
            theme = load_excel_theme(path)
            rule = first_matching_rule("submission_id", theme.column_formats)
            self.assertIsNotNone(rule)
            assert rule is not None
            self.assertTrue(match_column_rule("submission_id", rule))
        finally:
            path.unlink(missing_ok=True)

    def test_apply_partitioned_sheet_theme(self) -> None:
        theme = load_excel_theme(DEFAULT_PARTITIONED_THEME)
        wb = Workbook()
        ws = wb.active
        columns = [
            "area_id",
            "area_nombre",
            "periodo_mes_label",
        ]
        for col_idx, name in enumerate(columns, start=1):
            ws.cell(row=1, column=col_idx, value=name)
        ws.cell(row=2, column=1, value="trabajo_social")
        ws.cell(row=2, column=2, value="Trabajo Social")
        ws.cell(row=2, column=3, value="abr 2026")

        apply_sheet_theme(ws, wb, columns, theme)

        self.assertEqual(ws["A1"].value, "ID área")
        self.assertEqual(ws["B1"].value, "Área")
        self.assertTrue(ws.tables)


if __name__ == "__main__":
    unittest.main()
