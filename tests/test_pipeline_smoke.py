from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
LIB = PROJECT_ROOT / "lib"
if str(LIB) not in sys.path:
    sys.path.insert(0, str(LIB))

from vfiic_kpis.excel_export import (
    write_partitioned_workbook_for_form,
    write_stacked_comparativo_workbook,
)
from vfiic_kpis.io import read_form
from vfiic_kpis.manifest import reconcile
from vfiic_kpis.metrics import build_form_comparison
from vfiic_kpis.yaml_loader import load_forms_from_yaml


SCHEMA_BODY = """
Coordinacion Periciales:

  VFIIC KPIs Periciales - Demo:
    - columna_origen: "Visitas"
      descripcion: "Visitas realizadas"
    - columna_origen: "Dictamenes"
      descripcion: "Dictamenes realizados"
"""

DEMO_ROW = {
    "Periodo a Evaluar": "2026-04-01",
    "Agente - Nombre(s)": "Ana",
    "Agente - Apellido Paterno": "Lopez",
    "Agente - Apellido Materno": "Perez",
    "Visitas": 2,
    "Dictamenes": 1,
}


def _count_conditional_format_entries(ws) -> int:
    """Cuenta el total de reglas registradas (sumando todas las reglas de todos los rangos).

    Tras cargar el workbook con `load_workbook`, cada rango (sqref) puede tener
    una o varias reglas; aquí sumamos las longitudes de las listas de reglas
    para verificar que las tres reglas semánticas (>0, <0, =0) están presentes.
    """
    return sum(len(rules) for _, rules in ws.conditional_formatting._cf_rules.items())


def _build_pipeline(root: Path, df: pd.DataFrame):
    schema_path = root / "schema.yaml"
    schema_path.write_text(SCHEMA_BODY, encoding="utf-8")

    input_dir = root / "inputs"
    input_dir.mkdir()
    df.to_excel(
        input_dir / "VFIIC KPIs Periciales - Demo.xlsx",
        index=False,
    )

    forms = load_forms_from_yaml(schema_path)
    report = reconcile(forms, input_dir)
    assert len(report.matched) == 1
    resolution = report.matched[0]
    read_result = read_form(resolution)
    data = read_result.df
    comparison = build_form_comparison(data, resolution)
    assert comparison is not None
    return resolution, data, comparison


class TestPipelineSmokeNoYoy(unittest.TestCase):
    """Datos sólo con MoM: las columnas YoY se omiten globalmente."""

    def test_workbook_omits_yoy_columns_and_uses_conditional_format(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            df = pd.DataFrame(
                [
                    {**DEMO_ROW, "Periodo a Evaluar": "2026-03-01", "Visitas": 2, "Dictamenes": 1},
                    {**DEMO_ROW, "Periodo a Evaluar": "2026-04-01", "Visitas": 5, "Dictamenes": 3},
                    {
                        **DEMO_ROW,
                        "Periodo a Evaluar": "2026-04-01",
                        "Agente - Nombre(s)": "Bea",
                        "Visitas": 1,
                        "Dictamenes": 0,
                    },
                ]
            )
            resolution, data, comparison = _build_pipeline(root, df)

            stacked_out = root / "comparativo.xlsx"
            partitioned_out = root / "particionado.xlsx"

            write_stacked_comparativo_workbook(results=[comparison], output_path=stacked_out)
            write_partitioned_workbook_for_form(
                data=data,
                form_id=resolution.spec.area_id,
                output_path=partitioned_out,
            )

            self.assertTrue(stacked_out.exists())
            self.assertTrue(partitioned_out.exists())

            wb = load_workbook(stacked_out, data_only=False)
            ws = wb.active

            self.assertEqual(ws.max_column, 5)
            for header_col in range(6, 9):
                self.assertIsNone(ws.cell(row=2, column=header_col).value)

            first_kpi_row = 3
            second_kpi_row = 4
            d_mom = ws.cell(row=first_kpi_row, column=4)
            e_mom = ws.cell(row=first_kpi_row, column=5)
            self.assertEqual(d_mom.data_type, "f")
            self.assertEqual(e_mom.data_type, "f")
            for cell in (d_mom, e_mom):
                formula = str(cell.value)
                self.assertIn("ISNUMBER", formula)
                self.assertIn("B{0}".format(first_kpi_row), formula)
                self.assertIn("C{0}".format(first_kpi_row), formula)

            for col_idx in range(6, 9):
                self.assertIsNone(ws.cell(row=first_kpi_row, column=col_idx).value)

            self.assertTrue(e_mom.font.bold)
            b_row3 = ws.cell(row=first_kpi_row, column=2)
            b_row4 = ws.cell(row=second_kpi_row, column=2)
            self.assertEqual(b_row3.fill.fgColor.rgb, "FFFFFFFF")
            self.assertEqual(b_row4.fill.fgColor.rgb, "FFD9E2F3")

            self.assertGreaterEqual(_count_conditional_format_entries(ws), 3)


class TestPipelineSmokeWithYoy(unittest.TestCase):
    """Datos con historia hasta el mismo mes del año anterior: columnas YoY presentes."""

    def test_workbook_includes_yoy_columns_and_formulas(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            df = pd.DataFrame(
                [
                    {**DEMO_ROW, "Periodo a Evaluar": "2025-04-01", "Visitas": 4, "Dictamenes": 2},
                    {**DEMO_ROW, "Periodo a Evaluar": "2026-03-01", "Visitas": 2, "Dictamenes": 1},
                    {**DEMO_ROW, "Periodo a Evaluar": "2026-04-01", "Visitas": 5, "Dictamenes": 3},
                    {
                        **DEMO_ROW,
                        "Periodo a Evaluar": "2026-04-01",
                        "Agente - Nombre(s)": "Bea",
                        "Visitas": 1,
                        "Dictamenes": 0,
                    },
                ]
            )
            _resolution, _data, comparison = _build_pipeline(root, df)
            self.assertTrue(comparison.df["valor_anio_anterior"].notna().any())

            stacked_out = root / "comparativo.xlsx"
            write_stacked_comparativo_workbook(results=[comparison], output_path=stacked_out)
            self.assertTrue(stacked_out.exists())

            wb = load_workbook(stacked_out, data_only=False)
            ws = wb.active

            self.assertEqual(ws.max_column, 8)

            first_kpi_row = 3
            d_mom = ws.cell(row=first_kpi_row, column=4)
            e_mom = ws.cell(row=first_kpi_row, column=5)
            g_yoy = ws.cell(row=first_kpi_row, column=7)
            h_yoy = ws.cell(row=first_kpi_row, column=8)

            for cell in (d_mom, e_mom, g_yoy, h_yoy):
                self.assertEqual(cell.data_type, "f")
                self.assertIn("ISNUMBER", str(cell.value))

            self.assertIn("F{0}".format(first_kpi_row), str(g_yoy.value))
            self.assertIn("F{0}".format(first_kpi_row), str(h_yoy.value))

            self.assertTrue(e_mom.font.bold)
            self.assertTrue(h_yoy.font.bold)

            # Una regla por color (>0, <0, =0) en cada uno de los rangos MoM y YoY.
            self.assertGreaterEqual(_count_conditional_format_entries(ws), 6)


class TestPipelineSmokePartialRows(unittest.TestCase):
    """Filas inválidas intercaladas no deben impedir el comparativo."""

    def test_invalid_rows_are_skipped_but_valid_rows_generate_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            rows = [
                {
                    **DEMO_ROW,
                    "Periodo a Evaluar": "2026-03-01",
                    "Visitas": 2,
                    "Dictamenes": 1,
                },
                {
                    **DEMO_ROW,
                    "Periodo a Evaluar": pd.NaT,
                    "Agente - Nombre(s)": "",
                    "Agente - Apellido Paterno": "",
                    "Agente - Apellido Materno": "",
                    "Visitas": pd.NA,
                    "Dictamenes": pd.NA,
                },
                {
                    **DEMO_ROW,
                    "Periodo a Evaluar": "2026-04-01",
                    "Visitas": 5,
                    "Dictamenes": 3,
                },
                {
                    **DEMO_ROW,
                    "Periodo a Evaluar": "2026-06-28",
                    "Agente - Nombre(s)": "",
                    "Agente - Apellido Paterno": "",
                    "Agente - Apellido Materno": "",
                    "Visitas": pd.NA,
                    "Dictamenes": pd.NA,
                },
            ]
            df = pd.DataFrame(rows)
            resolution, data, comparison = _build_pipeline(root, df)
            self.assertEqual(len(data), 2)
            self.assertIsNotNone(comparison)
            self.assertGreaterEqual(len(comparison.df), 1)

            stacked_out = root / "comparativo.xlsx"
            write_stacked_comparativo_workbook(results=[comparison], output_path=stacked_out)
            self.assertTrue(stacked_out.exists())


if __name__ == "__main__":
    unittest.main()
