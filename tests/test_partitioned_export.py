from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
LIB = PROJECT_ROOT / "lib"
if str(LIB) not in sys.path:
    sys.path.insert(0, str(LIB))

from vfiic_kpis.excel_export import write_partitioned_workbook_for_form


class TestPartitionedExport(unittest.TestCase):
    def test_month_sheets_are_ordered_newest_first(self) -> None:
        data = pd.DataFrame(
            {
                "Periodo Evaluado": [
                    datetime(2026, 3, 1),
                    datetime(2026, 6, 1),
                    datetime(2026, 4, 1),
                ],
                "Visitas": [1, 3, 2],
                "periodo_dt": [
                    datetime(2026, 3, 1),
                    datetime(2026, 6, 1),
                    datetime(2026, 4, 1),
                ],
                "periodo_mes_key": ["2026_mar", "2026_jun", "2026_abr"],
                "_agent_sort": ["ana", "ana", "ana"],
            }
        )
        with tempfile.TemporaryDirectory() as tmp:
            out_path = Path(tmp) / "particionado.xlsx"
            write_partitioned_workbook_for_form(
                data=data,
                form_id="demo",
                output_path=out_path,
            )
            workbook = load_workbook(out_path)
            self.assertEqual(
                workbook.sheetnames,
                ["original", "2026_jun", "2026_abr", "2026_mar"],
            )


if __name__ == "__main__":
    unittest.main()
