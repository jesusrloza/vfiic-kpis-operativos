from __future__ import annotations

from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment

from openpyxl.utils import get_column_letter

from vfiic_kpis.excel_styling import (
    apply_semantic_conditional_format,
    apply_sheet_theme,
    finalize_uniform_widths,
    paint_block_body,
    paint_block_data_row_stripes,
    paint_block_header,
    paint_porcentaje_bold,
    paint_title_band,
)
from vfiic_kpis.excel_theme import load_excel_theme, resolve_header_label
from vfiic_kpis.metrics import FormComparisonResult
from vfiic_kpis.paths import DEFAULT_COMPARISON_THEME, DEFAULT_PARTITIONED_THEME

STACKED_COLUMNS_FULL: tuple[str, ...] = (
    "indicador",
    "valor_actual",
    "valor_mes_anterior",
    "diferencia_mom",
    "porcentaje_mom",
    "valor_anio_anterior",
    "diferencia_yoy",
    "porcentaje_yoy",
)

STACKED_COLUMNS_NO_YOY: tuple[str, ...] = (
    "indicador",
    "valor_actual",
    "valor_mes_anterior",
    "diferencia_mom",
    "porcentaje_mom",
)

# Compat: alias histórico (se mantenía como tuple inmutable de 8 columnas).
STACKED_COLUMNS: tuple[str, ...] = STACKED_COLUMNS_FULL

BLOCK_GAP_ROWS = 2

# Columnas derivadas en el comparativo apilado: fórmulas Excel (valores base editables por el usuario).
_FORMULA_COLUMN_NAMES: frozenset[str] = frozenset(
    {"diferencia_mom", "porcentaje_mom", "diferencia_yoy", "porcentaje_yoy"}
)


def _resolve_stacked_columns(results: list[FormComparisonResult]) -> tuple[str, ...]:
    """Modo "global any": las 3 columnas YoY se incluyen sólo si al menos un
    formulario alcanza 13 meses de historia (es decir, `valor_anio_anterior`
    no nulo en alguna fila). En caso contrario se omiten para todos los bloques.
    """
    has_any_yoy = any(
        result.df["valor_anio_anterior"].notna().any()
        for result in results
        if "valor_anio_anterior" in result.df.columns
    )
    return STACKED_COLUMNS_FULL if has_any_yoy else STACKED_COLUMNS_NO_YOY


def _comparativo_derived_formulas(row_idx: int, columns: tuple[str, ...]) -> dict[str, str]:
    """Fórmulas alineadas con `metrics._safe_percentage` (*100) y sin #DIV/0! / #VALUE!.

    Las referencias se calculan a partir de la posición real de cada columna
    base en `columns`, de modo que sigan apuntando bien aunque se omita el
    bloque YoY (sin columna F).
    """
    r = row_idx
    formulas: dict[str, str] = {}

    valor_actual_letter = get_column_letter(columns.index("valor_actual") + 1)
    valor_mes_anterior_letter = get_column_letter(columns.index("valor_mes_anterior") + 1)
    b = valor_actual_letter
    c = valor_mes_anterior_letter

    formulas["diferencia_mom"] = (
        f'=IF(OR(NOT(ISNUMBER({b}{r})),NOT(ISNUMBER({c}{r}))),"",{b}{r}-{c}{r})'
    )
    formulas["porcentaje_mom"] = (
        f'=IF(OR(NOT(ISNUMBER({b}{r})),NOT(ISNUMBER({c}{r}))),"",'
        f'IF({c}{r}=0,"",({b}{r}-{c}{r})/{c}{r}*100))'
    )

    if "valor_anio_anterior" in columns:
        f = get_column_letter(columns.index("valor_anio_anterior") + 1)
        formulas["diferencia_yoy"] = (
            f'=IF(OR(NOT(ISNUMBER({b}{r})),NOT(ISNUMBER({f}{r}))),"",{b}{r}-{f}{r})'
        )
        formulas["porcentaje_yoy"] = (
            f'=IF(OR(NOT(ISNUMBER({b}{r})),NOT(ISNUMBER({f}{r}))),"",'
            f'IF({f}{r}=0,"",({b}{r}-{f}{r})/{f}{r}*100))'
        )

    return formulas


def write_partitioned_workbook_for_form(
    *,
    data: pd.DataFrame,
    form_id: str,
    output_path: Path,
    theme_path: Path | None = None,
) -> None:
    """Genera un Excel particionado por mes para un solo formulario.

    Crea la hoja `original` con todos los registros y una hoja por cada mes con
    datos. Si no existe ningún `periodo_mes_key` parseable, no escribe nada.
    """
    if data.empty or "periodo_mes_key" not in data.columns:
        return
    months = sorted(data["periodo_mes_key"].dropna().unique())
    if not months:
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    theme = load_excel_theme(theme_path or DEFAULT_PARTITIONED_THEME)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        original = data.drop(columns=["_agent_sort"], errors="ignore")
        original.to_excel(writer, sheet_name="original", index=False)
        apply_sheet_theme(writer.book["original"], writer.book, list(original.columns), theme)

        for month_key in months:
            month_df = data[data["periodo_mes_key"] == month_key].copy()
            month_df = month_df.sort_values(by="_agent_sort", ascending=True)
            month_df = month_df.drop(columns=["_agent_sort"], errors="ignore")
            sheet_title = str(month_key)[:31]
            month_df.to_excel(writer, sheet_name=sheet_title, index=False)
            apply_sheet_theme(
                writer.book[sheet_title],
                writer.book,
                list(month_df.columns),
                theme,
            )

    # `form_id` se conserva en la firma para futuras anotaciones (metadata, propiedades).
    _ = form_id


def _resolve_block_labels(
    columns: Iterable[str],
    theme,
    *,
    mes_actual: str,
    mes_anterior: str | None,
    mes_anio_anterior: str | None,
) -> list[str]:
    """Mapea cada columna interna a su etiqueta visible para el bloque.

    Las etiquetas dependientes del mes (`valor_actual`, `valor_mes_anterior`,
    `valor_anio_anterior`) toman el mes correspondiente; el resto consulta el
    tema (`header_labels`).
    """
    labels: list[str] = []
    for internal in columns:
        if internal == "valor_actual":
            labels.append(mes_actual)
        elif internal == "valor_mes_anterior":
            labels.append(mes_anterior or "Mes anterior")
        elif internal == "valor_anio_anterior":
            labels.append(mes_anio_anterior or "Mismo mes año anterior")
        else:
            labels.append(resolve_header_label(internal, theme))
    return labels


def _build_block_row(row: pd.Series, columns: tuple[str, ...]) -> list:
    """Devuelve los valores en el orden de `columns`. `None` se mantiene para celdas vacías."""
    output: list = []
    for column in columns:
        if column not in row:
            output.append(None)
            continue
        value = row[column]
        if value is None:
            output.append(None)
        elif isinstance(value, float) and pd.isna(value):
            output.append(None)
        else:
            output.append(value)
    return output


def write_stacked_comparativo_workbook(
    results: list[FormComparisonResult],
    output_path: Path,
    theme_path: Path | None = None,
    sheet_name: str = "Comparativo",
) -> None:
    """Escribe un único workbook con un solo sheet apilado por formulario.

    Cada bloque consta de:
      - Título (merge horizontal con el `display_name` del formulario).
      - Encabezados con etiquetas dinámicas de mes y rótulos parentéticos.
      - Filas (una por KPI): valores base (actual, mes anterior, año anterior) y
        fórmulas Excel para diferencias y variaciones % (MoM y YoY).
      - Dos filas vacías como separador entre bloques.

    Las tres columnas YoY (`Mismo mes año anterior`, `Diferencia (año anterior)`
    y `Variación % (año anterior)`) se omiten globalmente cuando ningún
    formulario tiene 13 meses de historia ("global any"). El color de fuente
    en las celdas de diferencia/variación se aplica vía formato condicional
    Excel, por lo que se actualiza al recalcular fórmulas si el usuario edita
    los valores base. Al final se aplican anchos uniformes considerando todo
    el contenido.
    """
    if not results:
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    theme = load_excel_theme(theme_path or DEFAULT_COMPARISON_THEME)

    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet_name[:31] or "Comparativo"
    ws.sheet_view.showGridLines = theme.layout.show_grid_lines

    columns = _resolve_stacked_columns(results)
    column_count = len(columns)
    column_list = list(columns)

    diff_mom_idx = column_list.index("diferencia_mom") + 1
    pct_mom_idx = column_list.index("porcentaje_mom") + 1
    has_yoy_columns = "diferencia_yoy" in column_list
    diff_yoy_idx = column_list.index("diferencia_yoy") + 1 if has_yoy_columns else None
    pct_yoy_idx = column_list.index("porcentaje_yoy") + 1 if has_yoy_columns else None

    next_row = 1
    first_data_row: int | None = None
    semantic_ranges: list[str] = []
    porcentaje_indices = [pct_mom_idx]
    if pct_yoy_idx is not None:
        porcentaje_indices.append(pct_yoy_idx)

    for result in results:
        title_row = next_row
        header_row = title_row + 1
        data_row_start = header_row + 1

        paint_title_band(ws, result.spec.display_name, theme, row_idx=title_row, column_count=column_count)

        labels = _resolve_block_labels(
            column_list,
            theme,
            mes_actual=result.mes_actual_label,
            mes_anterior=result.mes_anterior_label,
            mes_anio_anterior=result.mes_anio_anterior_label,
        )
        paint_block_header(ws, theme, column_list, labels, row_idx=header_row)

        rows_written = 0
        for offset, (_, kpi_row) in enumerate(result.df.iterrows()):
            row_idx = data_row_start + offset
            block_values = _build_block_row(kpi_row, columns)
            derived = _comparativo_derived_formulas(row_idx, columns)
            for slot, internal in enumerate(columns, start=1):
                if internal in _FORMULA_COLUMN_NAMES:
                    cell = ws.cell(row=row_idx, column=slot, value=derived[internal])
                else:
                    cell = ws.cell(row=row_idx, column=slot, value=block_values[slot - 1])
                cell.alignment = Alignment(horizontal="left" if slot == 1 else "right", vertical="center")
            rows_written += 1

        if rows_written == 0:
            next_row = title_row + 1
            continue

        data_row_end = data_row_start + rows_written - 1

        paint_block_body(ws, theme, column_list, data_row_start=data_row_start, data_row_end=data_row_end)
        paint_block_data_row_stripes(
            ws,
            theme,
            data_row_start=data_row_start,
            data_row_end=data_row_end,
            column_count=column_count,
        )
        paint_porcentaje_bold(
            ws,
            column_indices=porcentaje_indices,
            data_row_start=data_row_start,
            data_row_end=data_row_end,
        )

        mom_start = get_column_letter(diff_mom_idx)
        mom_end = get_column_letter(pct_mom_idx)
        semantic_ranges.append(f"{mom_start}{data_row_start}:{mom_end}{data_row_end}")
        if diff_yoy_idx is not None and pct_yoy_idx is not None:
            yoy_start = get_column_letter(diff_yoy_idx)
            yoy_end = get_column_letter(pct_yoy_idx)
            semantic_ranges.append(f"{yoy_start}{data_row_start}:{yoy_end}{data_row_end}")

        if first_data_row is None:
            first_data_row = header_row

        next_row = data_row_end + 1 + BLOCK_GAP_ROWS

    apply_semantic_conditional_format(ws, theme, ranges=semantic_ranges)

    if first_data_row is not None and theme.layout.freeze_panes:
        ws.freeze_panes = ws.cell(row=first_data_row + 1, column=1).coordinate

    finalize_uniform_widths(ws, theme)
    workbook.save(output_path)
